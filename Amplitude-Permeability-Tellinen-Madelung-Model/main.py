"""Two-stage amplitude permeability experiment with Tellinen-Madelung memory.

Stage 1: begin at (+Hmajor, Upper(Hmajor)); perform AC demagnetization.
Stage 2: from the actual final state (including full reversal-stack memory),
         increase the oscillation amplitude to Hb, then measure the last cycle.
The same demagnetization history is computed once and reused for each Hb.

Place Major_Loop.csv beside the Python scripts and run: python main.py
"""
from dataclasses import dataclass, field
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference
from openpyxl.chart.series_factory import SeriesFactory as Series
from openpyxl.styles import Font

from Parameters import parameters as par
from BH import BH_increasing, BH_decreasing, Upper, H_MIN, H_MAX
from Hscan import Hdemag, Hscan, T, t_demag as T_DEMAG, Hswitch

mu0 = 4.0 * np.pi * 1e-7
N_DEMAG = par["DemagPeriods"]
N_REMAG = par["RemagPeriods"]
SAMPLES = par["SamplesPerPeriod"]
H_TOL = par["H_tol"]

# Exact period grid: shared endpoint is not computed twice.
demag_times = np.linspace(0.0, T_DEMAG, N_DEMAG*SAMPLES+1)
growth_times = T_DEMAG + np.arange(N_REMAG*SAMPLES+1)*T/SAMPLES


@dataclass
class State:
    H: float
    B: float
    root: str = "upper_major"
    stack: list = field(default_factory=list)
    direction: str = "dec"

    def clone(self):
        return State(self.H, self.B, self.root, list(self.stack), self.direction)


def propagate(fields, initial):
    """Walk through the field samples, carrying *all* hysteresis memory."""
    fields = np.asarray(fields, dtype=float)
    state = initial.clone()
    if abs(fields[0] - state.H) > H_TOL:
        raise ValueError("Discontinuous H between excitation stages")
    values = np.empty(len(fields), dtype=float)
    values[0] = state.B
    for i in range(1, len(fields)):
        H = float(fields[i])
        if abs(H-state.H) <= H_TOL:
            B = state.B
        else:
            if H > state.H:
                B, state.root, state.stack = BH_increasing(
                    H, state.direction == "inc", state.H, state.B,
                    state.root, state.stack)
                state.direction = "inc"
            else:
                B, state.root, state.stack = BH_decreasing(
                    H, state.direction == "dec", state.H, state.B,
                    state.root, state.stack)
                state.direction = "dec"
        state.H = H
        state.B = float(B)
        values[i] = B
    return values, state


def check_fields(fields):
    lo, hi = float(np.min(fields)), float(np.max(fields))
    if lo < H_MIN - H_TOL or hi > H_MAX + H_TOL:
        raise ValueError(
            f"Excitation interval [{lo:g},{hi:g}] outside major-loop "
            f"range [{H_MIN:g},{H_MAX:g}] A/m")


def prepare_demagnetized_state():
    """Do not replace the resulting state with H=B=0: keep actual memory."""
    Hd = np.asarray(Hdemag(demag_times), dtype=float)
    check_fields(Hd)
    init = State(float(Hd[0]), float(Upper(Hd[0])))
    Bd, state = propagate(Hd, init)
    return Hd, Bd, state


def calculate_B_history(Hb, prepared=None, include_demag=False):
    """Simulate growth for one target Hb, using a common demag prefix.

    Returns (H, B) for the growth stage by default; if include_demag=True,
    returns the complete two-stage trajectory without a duplicate endpoint.
    """
    if prepared is None:
        prepared = prepare_demagnetized_state()
    Hd, Bd, demag_state = prepared
    Hg = np.asarray(Hscan(growth_times, float(Hb)), dtype=float)
    check_fields(Hg)
    Bg, _ = propagate(Hg, demag_state)
    if include_demag:
        return np.concatenate([Hd[:-1], Hg]), np.concatenate([Bd[:-1], Bg])
    return Hg, Bg


def save_results_to_excel(results, filename, demag_state):
    wb = Workbook()
    ws = wb.active
    ws.title = "Amplitude permeability"
    headers = ["Hb (A/m)", "Bmax (T)", "mu_a", "Hmax last (A/m)",
               "Bmin (T)", "B at H=0 approx (T)", "B amplitude (T)",
               "mu_a from half-range"]
    for c, name in enumerate(headers, 1):
        ws.cell(1, c, name).font = Font(bold=True)
    for r, item in enumerate(results, 2):
        for c, key in enumerate(("Hb", "Bmax", "mu_a", "Hmax_last", "Bmin",
                                 "Bzero", "Bamp", "mu_a_range"), 1):
            ws.cell(r, c, float(item[key]))
    metadata = [
        ("Initial branch", "upper_major"), ("Hmajor (A/m)", par["Hmajor"]),
        ("Hfloor (A/m)", par["Hfloor"]), ("Hswitch (A/m)", Hswitch),
        ("Demag periods", N_DEMAG), ("Growth periods", N_REMAG),
        ("Samples/period", SAMPLES), ("B after demag (T)", demag_state.B),
        ("H after demag (A/m)", demag_state.H),
        ("Reversal stack entries", len(demag_state.stack)),
        ("mu0 (H/m)", mu0),
    ]
    ws["J1"] = "Parameter"
    ws["K1"] = "Value"
    ws["J1"].font = Font(bold=True)
    ws["K1"].font = Font(bold=True)
    for r, (key, value) in enumerate(metadata, 2):
        ws.cell(r, 10, key)
        ws.cell(r, 11, value)
    for column in "ABCDEFGH":
        ws.column_dimensions[column].width = 23
    ws.column_dimensions["J"].width = 27
    ws.column_dimensions["K"].width = 19
    ws.freeze_panes = "A2"

    chart = ScatterChart()
    chart.title = "Amplitude permeability after AC demagnetization"
    chart.x_axis.title = "Bmax (T)"
    chart.y_axis.title = "mu_a"
    chart.height, chart.width = 12, 20
    chart.legend = None
    xvalues = Reference(ws, min_col=2, min_row=2, max_row=len(results)+1)
    yvalues = Reference(ws, min_col=3, min_row=2, max_row=len(results)+1)
    series = Series(yvalues, xvalues, title="mu_a(B)")
    series.marker.symbol = "circle"
    chart.series.append(series)
    ws.add_chart(chart, "J15")
    wb.save(filename)


def main():
    prepared = prepare_demagnetized_state()
    Hd, Bd, demag_state = prepared
    print(f"After demagnetization: H={demag_state.H:.6g} A/m, "
          f"B={demag_state.B:.6g} T, stack={len(demag_state.stack)}")
    print("Note: the model does not force B=0 after AC demagnetization.")

    Hb_values = np.linspace(par["HbMin"], par["HbMax"], par["HbPoints"])
    results = []
    diagnostic_history = None
    for index, Hb in enumerate(Hb_values, 1):
        Hg, Bg = calculate_B_history(float(Hb), prepared)
        # Exactly one complete last period, including its endpoints.
        H_last = Hg[-(SAMPLES + 1):]
        B_last = Bg[-(SAMPLES + 1):]
        Bmax, Bmin = float(np.max(B_last)), float(np.min(B_last))
        Bamp = 0.5 * (Bmax - Bmin)
        Hmax = float(np.max(H_last))
        near_zero_idx = int(np.argmin(np.abs(H_last)))
        results.append(dict(Hb=float(Hb), Bmax=Bmax, Bmin=Bmin,
                            Bamp=Bamp, Bzero=float(B_last[near_zero_idx]),
                            mu_a=Bmax/(mu0*float(Hb)),
                            mu_a_range=Bamp/(mu0*float(Hb)),
                            Hmax_last=Hmax))
        print(f"{index:3d}/{len(Hb_values)}  Hb={Hb:9.3f} A/m  "
              f"Bmax={Bmax:10.6g} T  mu_a={results[-1]['mu_a']:.6g}")
        if index == len(Hb_values):
            diagnostic_history = (Hg, Bg)

    outdir = Path(__file__).resolve().parent
    save_results_to_excel(results, outdir/"Amplitude_Permeability.xlsx", demag_state)
    results.sort(key=lambda x: x["Bmax"])
    plt.figure(figsize=(7, 5))
    plt.plot([r["Bmax"] for r in results], [r["mu_a"] for r in results], "o-")
    plt.xlabel("Bmax (T)")
    plt.ylabel("Amplitude relative permeability, mu_a")
    plt.title("AC demagnetization + increasing-amplitude scan")
    plt.grid(True)
    plt.tight_layout()
    plt.savefig(outdir/"Amplitude_Permeability.png", dpi=160)

    # Record the full H(t) and the representative B(H) trajectory for checking.
    Hg, Bg = diagnostic_history
    plt.figure(figsize=(9, 4))
    t_full = np.concatenate([demag_times[:-1], growth_times])
    H_full = np.concatenate([Hd[:-1], Hg])
    plt.plot(t_full, H_full, linewidth=0.8)
    plt.axvline(T_DEMAG, ls="--", color="gray", label="End of demagnetization")
    plt.xlabel("Time (s)")
    plt.ylabel("H (A/m)")
    plt.title(f"Two-stage field excitation, Hb={Hb_values[-1]:g} A/m")
    plt.legend()
    plt.tight_layout()
    plt.savefig(outdir/"Excitation.png", dpi=160)

    plt.figure(figsize=(6, 6))
    plt.plot(Hd, Bd, linewidth=0.65, label="Demagnetization")
    plt.plot(Hg, Bg, linewidth=0.65, label="Amplitude growth")
    plt.xlabel("H (A/m)")
    plt.ylabel("B (T)")
    plt.title(f"B(H), Hb={Hb_values[-1]:g} A/m")
    plt.legend()
    plt.grid(True)
    plt.tight_layout()
    plt.savefig(outdir/"BH_Two_Stage.png", dpi=160)
    print(f"Saved xlsx and three diagnostic plots to {outdir}")
    plt.show()


if __name__ == "__main__":
    main()
